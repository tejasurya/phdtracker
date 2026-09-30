#!/usr/bin/env python3
"""
Converts data/tracker.xlsx into the JSON files the website reads.

  Scholarships → data/scholarships.json
  Positions    → data/excel_positions.json
  Professors   → data/professors.json
  Labs         → data/labs.json

'My Applications' and 'Document Checklist' are private and are never exported.
Headers are matched case-insensitively; row 1 must be the header row.
"""
import datetime as dt
import hashlib
import json
import pathlib
import re
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
XLSX = DATA / "tracker.xlsx"

SHEETS = {
    "Scholarships": ("scholarships.json", "scholarship"),
    "Positions": ("excel_positions.json", "title"),
    "Professors": ("professors.json", "name"),
    "Labs": ("labs.json", "lab_group"),
}


def key(header):
    return re.sub(r"[^a-z0-9]+", "_", str(header or "").lower().replace("*", "")).strip("_")


def cell(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return " ".join(v.split())
    return v


def split_areas(v):
    return [a.strip() for a in re.split(r"[;,]", v or "") if a.strip()] if isinstance(v, str) else []


def main():
    if not XLSX.exists():
        sys.exit(f"{XLSX} not found")
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    for sheet, (out, required) in SHEETS.items():
        if sheet not in wb.sheetnames:
            print(f"skip: sheet '{sheet}' missing")
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        headers = [key(h) for h in next(rows, [])]
        records = []
        for r in rows:
            rec = {h: cell(v) for h, v in zip(headers, r) if h}
            if not rec.get(required):
                continue
            rec = {k: v for k, v in rec.items() if v not in (None, "")}
            if "research_areas" in rec:
                rec["areas"] = split_areas(rec.pop("research_areas"))
            if "professors" in rec:
                rec["professors"] = split_areas(rec.pop("professors"))
            ident = f"{sheet}|{rec.get(required)}|{rec.get('link', '')}|{rec.get('university', rec.get('institution', ''))}"
            rec["id"] = sheet[0].lower() + "-" + hashlib.sha1(ident.encode()).hexdigest()[:10]
            records.append(rec)
        (DATA / out).write_text(json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{sheet}: {len(records)} rows → data/{out}")


if __name__ == "__main__":
    main()
